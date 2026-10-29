#!/usr/bin/env python3
"""Extend a rollout workbook with 300 non-duplicate pure primary schools."""

from __future__ import annotations

import argparse
import hashlib
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


DISTRICT_TARGET = 20
SIZE_BANDS = 5
PER_BAND = DISTRICT_TARGET // SIZE_BANDS


def stable_key(row: tuple) -> str:
    """Return the deterministic selection key used for the 2025/26 sample."""
    return hashlib.sha256(f"2025-26|{row[4]}|{row[0]}".encode()).hexdigest()


def school_number(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def assign_size_bands(rows: list[tuple]) -> list[tuple[tuple, int]]:
    ordered = sorted(rows, key=lambda row: (row[9], row[0]))
    count = len(ordered)
    return [
        (row, min(SIZE_BANDS - 1, index * SIZE_BANDS // count))
        for index, row in enumerate(ordered)
    ]


def select_schools(source_rows: list[tuple], existing: set[int]) -> list[tuple[tuple, int]]:
    by_district: dict[str, list[tuple]] = defaultdict(list)
    for row in source_rows:
        if str(row[2]).strip() == "G" and str(row[3]).strip() == "Grundschule":
            by_district[str(row[4]).strip()].append(row)

    if len(by_district) != 15:
        raise ValueError(f"Expected 15 school districts, found {len(by_district)}")

    selected: list[tuple[tuple, int]] = []
    for district in sorted(by_district):
        candidates: dict[int, list[tuple]] = defaultdict(list)
        for row, band in assign_size_bands(by_district[district]):
            if int(row[0]) not in existing:
                candidates[band].append(row)

        for band in range(SIZE_BANDS):
            available = sorted(candidates[band], key=stable_key)
            if len(available) < PER_BAND:
                raise ValueError(
                    f"District {district}, size band {band}: "
                    f"only {len(available)} candidates for quota {PER_BAND}"
                )
            selected.extend((row, band) for row in available[:PER_BAND])

    return selected


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path, help="All general-education schools 2025/26")
    parser.add_argument("input", type=Path, help="Rollout workbook to extend")
    parser.add_argument("output", type=Path, help="New extended workbook")
    args = parser.parse_args()

    source_ws = load_workbook(args.source, read_only=True, data_only=True)["Schulliste"]
    source_rows = [
        tuple(row)
        for row in source_ws.iter_rows(min_row=4, max_col=14, values_only=True)
        if row[0] is not None
    ]

    workbook = load_workbook(args.input)
    worksheet = workbook["Auswahl"]
    existing = {
        number
        for row in worksheet.iter_rows(values_only=True)
        if (number := school_number(row[1])) is not None
    }

    selected = select_schools(source_rows, existing)
    start_row = worksheet.max_row + 1
    for row, _band in selected:
        target_row = worksheet.max_row + 1
        worksheet.cell(target_row, 2, int(row[0]))
        worksheet.cell(target_row, 3, str(row[1]).strip())
        worksheet.cell(target_row, 6, str(row[4]).strip())

    filled_column_i = 0
    filled_column_j = 0
    for target_row in range(1, worksheet.max_row + 1):
        number = school_number(worksheet.cell(target_row, 2).value)
        if number is None:
            continue
        if worksheet.cell(target_row, 9).value is None:
            worksheet.cell(target_row, 9, f"poststelle{number}@schule.hessen.de")
            filled_column_i += 1
        if worksheet.cell(target_row, 10).value is None:
            worksheet.cell(
                target_row,
                10,
                f'=\"Schule(\"&B{target_row}&\")@schule.hessen.de\"',
            )
            filled_column_j += 1

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)

    added_numbers = [int(row[0]) for row, _band in selected]
    district_counts = Counter(str(row[4]).strip() for row, _band in selected)
    band_counts = Counter(band for _row, band in selected)
    if len(selected) != 300 or len(set(added_numbers)) != 300:
        raise AssertionError("Selection must contain 300 unique schools")
    if existing.intersection(added_numbers):
        raise AssertionError("Selection overlaps the existing rollout list")
    if set(district_counts.values()) != {DISTRICT_TARGET}:
        raise AssertionError(f"Unexpected district distribution: {district_counts}")
    if set(band_counts.values()) != {60}:
        raise AssertionError(f"Unexpected size-band distribution: {band_counts}")

    print(args.output)
    print(f"existing_rows={start_row - 1}")
    print(f"added_rows={len(selected)}")
    print(f"final_rows={worksheet.max_row}")
    print(f"filled_column_i={filled_column_i}")
    print(f"filled_column_j={filled_column_j}")
    print("districts=" + ",".join(f"{key}:{district_counts[key]}" for key in sorted(district_counts)))
    print("size_bands=" + ",".join(f"{key}:{band_counts[key]}" for key in sorted(band_counts)))


if __name__ == "__main__":
    main()
