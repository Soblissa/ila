#!/usr/bin/env python3
"""Erzeugt die Analyse des Barrierefreiheits-Prüfberichts als Markdown und DOCX."""

from __future__ import annotations

import argparse
import re
import subprocess
import tempfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REPORT_DATE = "19.10.2026"
TEST_PERIOD = "28.09.2026 bis 09.10.2026"
TARGET = "WCAG 2.2 AA und EN 301 549 Version 3.2.1 (Web)"


@dataclass
class Issue:
    number: int
    kind: str
    title: str
    severity: str
    trade: str
    criteria: list[str]
    category: str = ""


CATEGORY_IDS = {
    "Semantik, Struktur, Rollen und Beziehungen": {
        6, 8, 9, 17, 18, 36, 56, 58, 59, 61, 64, 65, 73, 78, 83, 90,
        93, 94, 95, 103, 114, 116,
    },
    "Tastatur, Fokus, Dialoge und Navigation": {
        10, 11, 12, 19, 20, 22, 24, 26, 27, 32, 34, 40, 42, 43, 44,
        80, 81, 82, 91, 101, 104, 118, 119, 120,
    },
    "Darstellung, Kontrast, Zoom und Benutzerpräferenzen": {
        21, 23, 25, 28, 29, 38, 39, 41, 46, 47, 48, 50, 53, 85, 86,
        92, 102, 117, 122,
    },
    "Formulare, Validierung und Statusmeldungen": {
        35, 45, 49, 52, 55, 69, 72, 74, 79, 87, 96, 99, 100, 105, 107,
        115, 123,
    },
    "Zugängliche Namen, Alternativtexte und Sprache": {
        7, 13, 14, 15, 30, 51, 54, 57, 66, 77, 84, 88, 97, 98, 121,
        124,
    },
}


SEVERITY_ORDER = {"Hoch": 0, "Mittel": 1, "Niedrig": 2, "Nicht bewertet": 3}
PRIORITY_LABEL = {
    "Hoch": "Priorität 1 – kritisch",
    "Mittel": "Priorität 2 – wesentlich",
    "Niedrig": "Priorität 3 – nachgelagert",
    "Nicht bewertet": "Priorität 4 – Empfehlungen ohne Schweregrad",
}


def extract_text(pdf: Path) -> str:
    with tempfile.TemporaryDirectory() as temp_dir:
        txt = Path(temp_dir) / "report.txt"
        subprocess.run(["pdftotext", "-layout", str(pdf), str(txt)], check=True)
        return txt.read_text(encoding="utf-8")


def parse_issues(text: str) -> list[Issue]:
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.strip() == "4 Identifizierte Probleme") + 1
    end = next(
        i for i, line in enumerate(lines[start:], start)
        if line.strip() == "5 Übersicht der Prüfschritte"
    )

    headings: list[tuple[int, int, re.Match[str]]] = []
    for i in range(start, end):
        current = lines[i].strip()
        if not re.match(r"^(Bug|Verstoß|Empfehlung):", current):
            continue
        parts = [current]
        j = i
        while not re.search(r"#\d+\s*$", " ".join(parts)) and j + 1 < end and j - i < 4:
            j += 1
            parts.append(lines[j].strip())
        heading = " ".join(part for part in parts if part)
        match = re.match(r"^(Bug|Verstoß|Empfehlung):\s*(.*?)\s*#(\d+)\s*$", heading)
        if match:
            headings.append((i, j, match))

    issues: list[Issue] = []
    for index, (_, heading_end, match) in enumerate(headings):
        next_start = headings[index + 1][0] if index + 1 < len(headings) else end
        body = "\n".join(lines[heading_end + 1:next_start])
        severity_match = re.search(r"Schwere des Problems\s+(Hoch|Mittel|Niedrig)", body)
        trade_match = re.search(r"Gewerk\s+([^\n]+)", body)
        criteria_block = re.search(
            r"Prüfschritte\s+(.*?)(?:Schwere des Problems|Gewerk)", body, re.S
        )
        criteria = []
        if criteria_block:
            criteria = re.findall(
                r"\b(?:5|6|7|9|11|12)\.\d+(?:\.\d+)*(?:[a-z])?\b",
                criteria_block.group(1),
            )
        issue = Issue(
            number=int(match.group(3)),
            kind=match.group(1),
            title=" ".join(match.group(2).split()),
            severity=severity_match.group(1) if severity_match else "Nicht bewertet",
            trade=trade_match.group(1).strip() if trade_match else "Nicht angegeben",
            criteria=criteria,
        )
        issues.append(issue)

    if len(issues) != 98:
        raise ValueError(f"Erwartet wurden 98 Feststellungen, gefunden wurden {len(issues)}")

    assigned: dict[int, str] = {}
    for category, ids in CATEGORY_IDS.items():
        for issue_id in ids:
            if issue_id in assigned:
                raise ValueError(f"Issue #{issue_id} ist doppelt kategorisiert")
            assigned[issue_id] = category
    parsed_ids = {issue.number for issue in issues}
    if set(assigned) != parsed_ids:
        raise ValueError(
            f"Kategorisierung unvollständig: fehlt={sorted(parsed_ids-set(assigned))}, "
            f"zusätzlich={sorted(set(assigned)-parsed_ids)}"
        )
    for issue in issues:
        issue.category = assigned[issue.number]
    return issues


def normalize_title(title: str) -> str:
    return (
        title.replace("Unnoetiger", "Unnötiger")
        .replace("Toogles", "Toggles")
        .replace("Fehlerfelder", "Fehlerfeldern")
    )


def ids(items: list[Issue]) -> str:
    return ", ".join(f"#{item.number}" for item in items)


def make_markdown(issues: list[Issue], source_name: str) -> str:
    kind_counts = Counter(issue.kind for issue in issues)
    severity_counts = Counter(issue.severity for issue in issues)
    lines = [
        "# Analyse des Prüfberichts zur Barrierefreiheit von diSAF",
        "",
        f"**Stand:** {REPORT_DATE}  ",
        f"**Prüfzeitraum:** {TEST_PERIOD}  ",
        f"**Konformitätsziel:** {TARGET}  ",
        f"**Quelle:** `{source_name}` (83 Seiten)",
        "",
        "## 1. Zusammenfassung",
        "",
        "Der Prüfbericht weist für keinen der zwölf getesteten Dialoge Barrierefreiheit aus. "
        "Festgestellt wurden Barrieren für blinde, sehbehinderte, motorisch und kognitiv "
        "eingeschränkte Nutzende. Für Menschen mit Hörbeeinträchtigungen wurden in der "
        "geprüften Seitenauswahl keine Barrieren festgestellt.",
        "",
        f"Insgesamt enthält der Bericht **{len(issues)} Feststellungen**: "
        f"{kind_counts['Verstoß']} Verstöße, {kind_counts['Bug']} Bugs und "
        f"{kind_counts['Empfehlung']} Empfehlungen. {severity_counts['Hoch']} sind hoch, "
        f"{severity_counts['Mittel']} mittel und {severity_counts['Niedrig']} niedrig "
        f"gewichtet; {severity_counts['Nicht bewertet']} Empfehlungen haben keinen "
        "ausgewiesenen Schweregrad.",
        "",
        "Die kritischsten Risiken liegen bei Screenreader-Semantik, zugänglichen Namen, "
        "Fokus- und Dialogsteuerung sowie nicht bedienbaren interaktiven Komponenten. "
        "Da sich viele Befunde auf wiederverwendete Komponenten beziehen, sollte die "
        "Behebung komponentenweit erfolgen und anschließend in allen zwölf geprüften "
        "Dialogen regressionsgetestet werden.",
        "",
        "### Berichtseigenheiten",
        "",
        "- Die Kennzahl „Gefundene Fehler: 72“ entspricht der Zahl der als „Verstoß“ "
        "  bezeichneten Punkte. Zusätzlich sind 12 als „Bug“ und 14 als „Empfehlung“ "
        "  bezeichnete Feststellungen enthalten.",
        "- Die Gewichtung 14 hoch + 65 mittel + 8 niedrig umfasst 87 Feststellungen. "
        "  Elf Empfehlungen besitzen keinen Schweregrad.",
        "- In der Detailauswertung ist #94 dem Gewerk „Struktur & Semantik“ zugeordnet; "
        "  die Management-Statistik führt dieses Gewerk nicht separat auf. #40 ist "
        "  zugleich Entwicklung und Design zugeordnet.",
        "",
        "## 2. Sortierung nach Art",
        "",
        "| Art | Anzahl | Bedeutung |",
        "|---|---:|---|",
        f"| Verstoß | {kind_counts['Verstoß']} | Muss für die angestrebte Konformität behoben werden. |",
        f"| Bug | {kind_counts['Bug']} | Technischer Fehler; in der Detailauswertung ebenfalls gewichtet. |",
        f"| Empfehlung | {kind_counts['Empfehlung']} | Verbessert Zugänglichkeit bzw. Bedienkomfort; 3 davon sind mittel gewichtet. |",
        "",
        "## 3. Sortierung nach Inhalt",
        "",
        "| Inhaltsbereich | Anzahl | Zugehörige Issues |",
        "|---|---:|---|",
    ]
    for category in CATEGORY_IDS:
        items = [issue for issue in issues if issue.category == category]
        lines.append(f"| {category} | {len(items)} | {ids(items)} |")

    lines.extend([
        "",
        "## 4. Priorisierung nach Schweregrad",
        "",
        "Die Priorisierung übernimmt den Schweregrad des Prüfberichts. Innerhalb einer "
        "Stufe sind die Feststellungen nach Issue-Nummer sortiert. Auch niedrige Verstöße "
        "müssen für WCAG-Konformität behoben werden.",
        "",
    ])
    for severity in ("Hoch", "Mittel", "Niedrig", "Nicht bewertet"):
        items = sorted((issue for issue in issues if issue.severity == severity), key=lambda x: x.number)
        lines.append(f"### {PRIORITY_LABEL[severity]} ({len(items)})")
        lines.append("")
        lines.append("| Issue | Art | Feststellung | Gewerk | Inhaltsbereich |")
        lines.append("|---:|---|---|---|---|")
        for issue in items:
            lines.append(
                f"| #{issue.number} | {issue.kind} | {normalize_title(issue.title)} | "
                f"{issue.trade} | {issue.category} |"
            )
        lines.append("")

    lines.extend([
        "## 5. Empfohlene Umsetzungsreihenfolge",
        "",
        "1. **Kritische Barrieren schließen:** alle 14 hoch gewichteten Feststellungen "
        "   bearbeiten und separat mit Tastatur sowie JAWS/NVDA prüfen.",
        "2. **Basiskomponenten sanieren:** Navigation, Dialog, Schalter/Switch, Aufklapper, "
        "   Link, Listen-/Kachelstruktur und zugängliche Beschriftungen zentral korrigieren.",
        "3. **Formulare und Rückmeldungen konsolidieren:** Fehlerzustände, Fehlermeldungen, "
        "   Statusmeldungen und Feldbeschriftungen als einheitliches Muster umsetzen.",
        "4. **Visuelle Barrieren beheben:** Fokusindikatoren, Text-/Grafikkontraste, "
        "   Hochkontrastmodus, Vergrößerung und Umbruch prüfen.",
        "5. **Abschlussprüfung durchführen:** alle zwölf Dialoge und übertragbare Vorkommen "
        "   außerhalb der Stichprobe erneut testen; erst danach Konformität bewerten.",
        "",
        "## 6. Kritische Feststellungen im Überblick",
        "",
    ])
    for issue in sorted((x for x in issues if x.severity == "Hoch"), key=lambda x: x.number):
        lines.append(
            f"- **#{issue.number} – {normalize_title(issue.title)}** "
            f"({issue.kind}; {issue.trade}; {issue.category})"
        )
    lines.append("")
    return "\n".join(lines)


BLUE = "003E70"
LIGHT_BLUE = "D9EAF4"
MID_BLUE = "2B6F9F"
LIGHT_GRAY = "F2F4F6"
RED = "A61B1B"
ORANGE = "B85C00"
GREEN = "527A24"


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, header=True, font_size=8.5) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if header and row_index == 0:
                shade_cell(cell, BLUE)
            elif row_index % 2 == 0:
                shade_cell(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    if header and row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = value
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    style_table(table)
    return table


def add_page_break(document: Document) -> None:
    document.add_page_break()


def add_bullet(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def make_docx(issues: list[Issue], output: Path, source_name: str) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color in (
        ("Title", 28, BLUE),
        ("Heading 1", 18, BLUE),
        ("Heading 2", 13, MID_BLUE),
        ("Heading 3", 11, BLUE),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    # Kopf- und Fußzeile
    header = section.header.paragraphs[0]
    header.text = "diSAF · Analyse des Prüfberichts zur Barrierefreiheit"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID_BLUE)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Stand: 19.10.2026  ·  Seite ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)

    # Titelblatt
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.add_run("ANALYSE").bold = True
    p.runs[0].font.size = Pt(13)
    p.runs[0].font.color.rgb = RGBColor.from_string(MID_BLUE)
    title = document.add_paragraph(style="Title")
    title.add_run("Prüfbericht zur\nBarrierefreiheit von diSAF")
    subtitle = document.add_paragraph()
    subtitle.add_run("Zusammenfassung · thematische Sortierung · Priorisierung").bold = True
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MID_BLUE)
    document.add_paragraph("")
    meta = add_table(
        document,
        ["Merkmal", "Angabe"],
        [
            ["Auftraggeber", "Hessisches Ministerium für Kultus, Bildung und Chancen"],
            ["Prüfzeitraum", TEST_PERIOD],
            ["Konformitätsziel", TARGET],
            ["Geprüfter Umfang", "12 Dialoge / Seitenbereiche"],
            ["Quelle", f"{source_name} (83 Seiten)"],
            ["Analysestand", REPORT_DATE],
        ],
        [4.0, 12.5],
    )
    document.add_paragraph("")
    note = document.add_paragraph()
    note.add_run("Ergebnis in einem Satz: ").bold = True
    note.add_run(
        "Keiner der zwölf geprüften Dialoge ist barrierefrei; die 98 Feststellungen "
        "erfordern eine komponentenweite, nach Schweregrad gesteuerte Nachbesserung."
    )

    add_page_break(document)

    kind_counts = Counter(issue.kind for issue in issues)
    severity_counts = Counter(issue.severity for issue in issues)

    document.add_heading("1. Zusammenfassung", level=1)
    document.add_paragraph(
        "Der Prüfbericht bewertet die geprüfte diSAF-Webanwendung gegen WCAG 2.2 auf "
        "Konformitätsstufe AA und EN 301 549. Nach dem Prüfbericht kann für keinen der "
        "zwölf geprüften Dialoge Barrierefreiheit ausgewiesen werden. Die festgestellten "
        "Barrieren betreffen blinde, sehbehinderte, motorisch und kognitiv eingeschränkte "
        "Nutzende. Für Menschen mit Hörbeeinträchtigungen wurden in der geprüften "
        "Seitenauswahl keine Barrieren festgestellt."
    )
    document.add_paragraph(
        "Die Befunde verteilen sich über alle vier WCAG-Grundsätze: Wahrnehmbarkeit, "
        "Bedienbarkeit, Verständlichkeit und Robustheit. Besonders kritisch sind "
        "Screenreader-Semantik, zugängliche Beschriftungen, Fokus- und Dialogsteuerung "
        "sowie nicht bedienbare interaktive Komponenten."
    )
    add_table(
        document,
        ["Kennzahl", "Wert", "Einordnung"],
        [
            ["Feststellungen gesamt", "98", "72 Verstöße + 12 Bugs + 14 Empfehlungen"],
            ["Hoch", str(severity_counts["Hoch"]), "kritische Barrieren mit Ausschlussrisiko"],
            ["Mittel", str(severity_counts["Mittel"]), "wesentliche Barrieren / Standardgewichtung"],
            ["Niedrig", str(severity_counts["Niedrig"]), "geringere Beeinträchtigung, dennoch konformitätsrelevant"],
            ["Ohne Schweregrad", str(severity_counts["Nicht bewertet"]), "Empfehlungen des Prüfteams"],
            ["Barrierefreie Dialoge", "0 von 12", "keine Konformität in der Stichprobe"],
        ],
        [5.0, 2.5, 9.0],
    )

    document.add_heading("1.1 Kernaussagen für die Steuerung", level=2)
    add_bullet(document, "Die 14 hoch gewichteten Feststellungen sind zuerst zu schließen.")
    add_bullet(document, "Viele Probleme liegen in wiederverwendeten Komponenten; punktuelle Einzelfixes reichen daher nicht aus.")
    add_bullet(document, "Entwicklung ist das am stärksten betroffene Gewerk, gefolgt von Design und Redaktion.")
    add_bullet(document, "Auch niedrig gewichtete Verstöße müssen für die angestrebte Konformität behoben werden.")
    add_bullet(document, "Nach der Umsetzung ist eine vollständige Regressionsprüfung aller zwölf Dialoge mit Tastatur und Screenreader erforderlich.")

    document.add_heading("1.2 Berichtseigenheiten und Abgrenzung", level=2)
    add_bullet(
        document,
        "Die Management-Kennzahl „Gefundene Fehler: 72“ entspricht exakt der Zahl der als „Verstoß“ bezeichneten Punkte. Zusätzlich enthält der Bericht 12 Bugs und 14 Empfehlungen.",
    )
    add_bullet(
        document,
        "Die Gewichtung 14 hoch + 65 mittel + 8 niedrig umfasst 87 Feststellungen. Elf Empfehlungen besitzen keinen ausgewiesenen Schweregrad.",
    )
    add_bullet(
        document,
        "Issue #94 ist in der Detailauswertung dem Gewerk „Struktur & Semantik“ zugeordnet, das in der Management-Statistik nicht separat erscheint. Issue #40 nennt zugleich Entwicklung und Design.",
    )
    add_bullet(
        document,
        "Geprüft wurde eine Seitenauswahl. Wiederkehrende Komponenten und vergleichbare Vorkommen außerhalb der Stichprobe müssen mitkorrigiert und nachgetestet werden.",
    )

    add_page_break(document)
    document.add_heading("2. Sortierung der Feststellungen nach Art", level=1)
    add_table(
        document,
        ["Art", "Anzahl", "Anteil", "Bedeutung"],
        [
            ["Verstoß", str(kind_counts["Verstoß"]), f"{kind_counts['Verstoß']/98:.0%}", "Für die Konformität zwingend zu beheben"],
            ["Bug", str(kind_counts["Bug"]), f"{kind_counts['Bug']/98:.0%}", "Technischer Fehler; im Bericht gewichtet"],
            ["Empfehlung", str(kind_counts["Empfehlung"]), f"{kind_counts['Empfehlung']/98:.0%}", "Verbesserung; 3 Empfehlungen sind mittel gewichtet"],
        ],
        [3.0, 2.0, 2.0, 9.5],
    )
    document.add_paragraph(
        "Für die Umsetzungssteuerung sollten Verstöße und Bugs gemeinsam als "
        "konformitätsrelevanter Arbeitsvorrat behandelt werden. Empfehlungen werden "
        "separat geplant, außer wenn der Bericht ihnen ausdrücklich den Schweregrad "
        "„Mittel“ zuweist (#9, #17 und #36)."
    )

    document.add_heading("3. Sortierung nach Inhalt", level=1)
    rows = []
    for category in CATEGORY_IDS:
        items = [issue for issue in issues if issue.category == category]
        rows.append([category, str(len(items)), ids(items)])
    add_table(document, ["Inhaltsbereich", "Anzahl", "Issues"], rows, [6.0, 1.8, 8.7])

    category_summaries = {
        "Semantik, Struktur, Rollen und Beziehungen":
            "Falsche oder fehlende Rollen, Listen-, Überschriften- und Linkstrukturen sowie inkonsistente ARIA-Beziehungen verhindern eine verlässliche Interpretation durch assistive Technologien.",
        "Tastatur, Fokus, Dialoge und Navigation":
            "Fokusverluste, Fokusüberlagerungen, Tastaturfallen und fehlerhafte Dialogsteuerung erschweren oder verhindern die Bedienung ohne Maus.",
        "Darstellung, Kontrast, Zoom und Benutzerpräferenzen":
            "Unzureichende Kontraste, Probleme bei Vergrößerung und Umbruch sowie Lücken im Hochkontrastmodus betreffen vor allem sehbehinderte Nutzende.",
        "Formulare, Validierung und Statusmeldungen":
            "Feldbeschriftungen, Fehlerzustände und Erfolgsmeldungen sind teilweise nicht programmatisch ermittelbar oder werden vom Screenreader nicht angesagt.",
        "Zugängliche Namen, Alternativtexte und Sprache":
            "Mehrere Schalter, Grafiken und Statussymbole besitzen keine, unklare oder inkonsistente zugängliche Beschriftungen.",
    }
    for category in CATEGORY_IDS:
        document.add_heading(category, level=2)
        document.add_paragraph(category_summaries[category])

    add_page_break(document)
    document.add_heading("4. Priorisierung nach Schweregrad", level=1)
    document.add_paragraph(
        "Die folgende Liste übernimmt die Gewichtung des Prüfberichts. Innerhalb jeder "
        "Stufe sind die Feststellungen nach Issue-Nummer sortiert. Die Prioritätsstufen "
        "ersetzen nicht die WCAG-Logik „bestanden / nicht bestanden“: Auch niedrig "
        "gewichtete Verstöße bleiben für die Konformität relevant."
    )
    for severity in ("Hoch", "Mittel", "Niedrig", "Nicht bewertet"):
        items = sorted((issue for issue in issues if issue.severity == severity), key=lambda x: x.number)
        document.add_heading(f"{PRIORITY_LABEL[severity]} ({len(items)})", level=2)
        table = add_table(
            document,
            ["Issue", "Art", "Feststellung", "Gewerk", "Inhaltsbereich"],
            [[
                f"#{issue.number}", issue.kind, normalize_title(issue.title), issue.trade, issue.category
            ] for issue in items],
            [1.3, 2.0, 6.4, 2.6, 4.2],
        )
        if severity == "Hoch":
            for row in table.rows[1:]:
                shade_cell(row.cells[0], "F4CCCC")
                row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(RED)
                row.cells[0].paragraphs[0].runs[0].font.bold = True

    add_page_break(document)
    document.add_heading("5. Empfohlene Umsetzungsreihenfolge", level=1)
    phases = [
        ("1", "Kritische Barrieren schließen", "Alle 14 hoch gewichteten Feststellungen bearbeiten. Abnahmekriterium: Funktionsprüfung mit Tastatur sowie JAWS und NVDA."),
        ("2", "Basiskomponenten sanieren", "Navigation, Dialog, Schalter/Switch, Aufklapper, Links, Listen/Kacheln und zugängliche Namen zentral korrigieren. Alle Vorkommen der Komponenten einbeziehen."),
        ("3", "Formulare und Rückmeldungen konsolidieren", "Einheitliche Muster für Feldbeschriftungen, Fehlerzustände, Fehlermeldungen, Pflichtfelder und Statusmeldungen umsetzen."),
        ("4", "Visuelle Barrieren beheben", "Fokusindikatoren, Text-/Grafikkontraste, Hochkontrastmodus, 200-Prozent-Vergrößerung und Umbruch prüfen und korrigieren."),
        ("5", "Regressions- und Konformitätsprüfung", "Alle zwölf Dialoge sowie übertragbare Vorkommen außerhalb der Stichprobe erneut prüfen; offene Empfehlungen dokumentiert entscheiden."),
    ]
    add_table(document, ["Phase", "Arbeitspaket", "Ziel / Abnahmekriterium"], [list(x) for x in phases], [1.4, 4.6, 10.5])

    document.add_heading("5.1 Kritische Feststellungen", level=2)
    for issue in sorted((x for x in issues if x.severity == "Hoch"), key=lambda x: x.number):
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(f"#{issue.number} – {normalize_title(issue.title)}").bold = True
        paragraph.add_run(f" ({issue.kind}; {issue.trade})")

    document.add_heading("6. Schlussfolgerung", level=1)
    document.add_paragraph(
        "Die aktuelle Fassung der geprüften diSAF-Anwendung erfüllt das angestrebte "
        "Konformitätsziel nicht. Der Schwerpunkt liegt nicht auf einzelnen kosmetischen "
        "Korrekturen, sondern auf systematischen Mängeln in Semantik, Komponentenlogik, "
        "Fokusführung und Rückmeldungen. Eine komponentenorientierte Sanierung ist daher "
        "effizienter und verlässlicher als die isolierte Bearbeitung einzelner Seiten."
    )
    document.add_paragraph(
        "Nach Abschluss der priorisierten Behebung sollte eine erneute Prüfung unter "
        "denselben Rahmenbedingungen erfolgen. Erst ein erfolgreicher Re-Test aller "
        "konformitätsrelevanten Kriterien kann die Grundlage für eine Aussage zur "
        "Barrierefreiheit bilden."
    )

    document.core_properties.title = "Analyse des Prüfberichts zur Barrierefreiheit von diSAF"
    document.core_properties.subject = "Zusammenfassung, Sortierung und Priorisierung der Feststellungen"
    document.core_properties.author = "Cheko – ILA-Hauptagent"
    document.core_properties.keywords = "diSAF, Barrierefreiheit, WCAG 2.2, EN 301 549"
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pdf", type=Path)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--markdown", type=Path, required=True)
    args = parser.parse_args()

    text = extract_text(args.pdf)
    issues = parse_issues(text)
    args.docx.parent.mkdir(parents=True, exist_ok=True)
    args.markdown.parent.mkdir(parents=True, exist_ok=True)
    args.markdown.write_text(make_markdown(issues, args.pdf.name), encoding="utf-8")
    make_docx(issues, args.docx, args.pdf.name)

    print(f"DOCX: {args.docx}")
    print(f"Markdown: {args.markdown}")
    print(f"Feststellungen: {len(issues)}")
    print(f"Arten: {dict(Counter(issue.kind for issue in issues))}")
    print(f"Schweregrade: {dict(Counter(issue.severity for issue in issues))}")


if __name__ == "__main__":
    main()
